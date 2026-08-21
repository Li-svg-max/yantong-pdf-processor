from __future__ import annotations

import base64
import unittest
from io import BytesIO

from docx import Document
from pydantic import ValidationError
from pypdf import PdfReader

from app.document_export.models import (
    ClozeExportRequest,
    ExportRequest,
    SignedClozeExportRequest,
)
from app.document_export.renderer import render_cloze_export, render_export
from app.request_auth import sign_ticket, ticket_validation_error


def build_request(format_name: str, package_mode: str = "combined") -> ExportRequest:
    return ExportRequest(
        format=format_name,
        title="数学二错题集",
        settings={
            "paper": "A4",
            "columns": 2,
            "fontSize": 11,
            "blankLines": 1,
            "packageMode": package_mode,
            "solutionMode": "end",
            "includeAnalysis": True,
        },
        groups=[{
            "order": 1,
            "title": "函数极限",
            "meta": "数学二 · 极限与连续",
            "questions": [{
                "number": "1",
                "stem": "求函数在 x=0 处的极限。",
                "options": ["A. 0", "B. 1"],
                "answer": "B",
                "analysis": "使用等价无穷小。",
            }],
        }],
    )


class DocumentExportTests(unittest.TestCase):
    def test_pdf_is_readable(self) -> None:
        files = render_export(build_request("pdf"))
        reader = PdfReader(BytesIO(base64.b64decode(files[0]["contentBase64"])))
        self.assertGreaterEqual(len(reader.pages), 2)

    def test_separate_docx_returns_two_readable_files(self) -> None:
        files = render_export(build_request("docx", "separate"))
        self.assertEqual(len(files), 2)
        for item in files:
            self.assertTrue(Document(BytesIO(base64.b64decode(item["contentBase64"]))).paragraphs)

    def test_cloze_docx_contains_exercise_and_answer_table(self) -> None:
        request = ClozeExportRequest(
            title="数据结构第三章",
            courseName="数据结构",
            blocks=[{
                "id": "block-1",
                "page": 2,
                "order": 1,
                "text": "线性表是由有限个数据元素组成的序列。",
                "blanks": [{
                    "id": "blank-1",
                    "start": 0,
                    "end": 3,
                    "answer": "线性表",
                    "enabled": True,
                }],
            }],
        )
        files = render_cloze_export(request)
        self.assertEqual(len(files), 1)
        document = Document(BytesIO(base64.b64decode(files[0]["contentBase64"])))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("〔1〕", text)
        self.assertIn("答案与核对", text)
        self.assertEqual(document.tables[0].cell(1, 1).text, "线性表")

    def test_cloze_export_requires_a_short_lived_unchanged_ticket(self) -> None:
        request = ClozeExportRequest(
            title="操作系统复习",
            courseName="操作系统",
            blocks=[{
                "id": "block-1",
                "page": 1,
                "order": 1,
                "text": "进程是资源分配的基本单位。",
                "blanks": [{
                    "id": "blank-1",
                    "start": 0,
                    "end": 2,
                    "answer": "进程",
                    "enabled": True,
                }],
            }],
        )
        now_ms = 2_000_000_000_000
        expires_at = now_ms + 300_000
        token = "test-processor-token"
        signature = sign_ticket(
            request,
            expires_at,
            token,
            payload_key="request",
        )
        self.assertEqual(
            signature,
            "2e86f75fd2e6bda8db16fb10961301c0f75058177df135b41e75350defda038e",
        )

        self.assertIsNone(
            ticket_validation_error(
                request,
                expires_at,
                signature,
                token,
                now_ms,
                payload_key="request",
            )
        )
        self.assertEqual(
            ticket_validation_error(
                request,
                now_ms - 1,
                signature,
                token,
                now_ms,
                payload_key="request",
            ),
            "job ticket expired",
        )
        modified = request.model_copy(update={"title": "被篡改的标题"})
        self.assertEqual(
            ticket_validation_error(
                modified,
                expires_at,
                signature,
                token,
                now_ms,
                payload_key="request",
            ),
            "job ticket signature does not match",
        )
        with self.assertRaises(ValidationError):
            SignedClozeExportRequest(
                request=request,
                expiresAt=expires_at,
                signature="",
            )


if __name__ == "__main__":
    unittest.main()
