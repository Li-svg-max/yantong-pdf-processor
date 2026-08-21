from __future__ import annotations

import base64
import html
import ipaddress
import os
import re
import socket
import urllib.request
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image as PillowImage
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, A5
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    HRFlowable,
    Image,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
)

from .models import (
    ClozeExportRequest,
    ExportGroup,
    ExportImage,
    ExportQuestion,
    ExportRequest,
)


MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_TOTAL_IMAGE_BYTES = 40 * 1024 * 1024
MAX_OUTPUT_BYTES = 12 * 1024 * 1024


def _safe_name(value: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|\r\n]+", "-", value).strip(" .-")
    return cleaned[:70] or "错题集"


def _public_host(hostname: str) -> bool:
    try:
        addresses = socket.getaddrinfo(hostname, 443, type=socket.SOCK_STREAM)
    except socket.gaierror:
        return False
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            return False
    return True


def _download_image(image: ExportImage) -> bytes:
    parsed = urlparse(image.url)
    if parsed.scheme != "https" or not parsed.hostname or not _public_host(parsed.hostname):
        raise ValueError("题目图片地址不是可访问的 HTTPS 云存储地址")
    request = urllib.request.Request(image.url, headers={"User-Agent": "YantongDocumentExport/1.0"})
    with urllib.request.urlopen(request, timeout=15) as response:
        content_type = str(response.headers.get("Content-Type", ""))
        if content_type and not content_type.lower().startswith("image/"):
            raise ValueError("云存储地址返回的不是图片")
        length = int(response.headers.get("Content-Length") or 0)
        if length > MAX_IMAGE_BYTES:
            raise ValueError("单张题目图片超过 8MB")
        data = response.read(MAX_IMAGE_BYTES + 1)
    if not data or len(data) > MAX_IMAGE_BYTES:
        raise ValueError("单张题目图片超过 8MB 或内容为空")
    with PillowImage.open(BytesIO(data)) as source:
        source.verify()
    return data


class ImageCache:
    def __init__(self) -> None:
        self._items: dict[str, bytes] = {}
        self._total = 0

    def get(self, image: ExportImage) -> bytes:
        if image.url not in self._items:
            data = _download_image(image)
            self._total += len(data)
            if self._total > MAX_TOTAL_IMAGE_BYTES:
                raise ValueError("本次导出的图片总量超过 40MB，请拆分错题篮后重试")
            self._items[image.url] = data
        return self._items[image.url]


def _document_modes(request: ExportRequest) -> list[str]:
    return ["questions", "solutions"] if request.settings.packageMode == "separate" else ["combined"]


def _suffix(mode: str) -> str:
    return {"questions": "-试题", "solutions": "-答案解析", "combined": ""}[mode]


def _question_label(group: ExportGroup, question: ExportQuestion, index: int) -> str:
    return question.number or (str(group.order) if index == 0 else f"{group.order}.{index + 1}")


def _pdf_paragraph(value: str, style: ParagraphStyle) -> Paragraph:
    text = html.escape(str(value or "")).replace("\n", "<br/>") or " "
    return Paragraph(text, style)


def _pdf_image(data: bytes, max_width: float) -> Image:
    with PillowImage.open(BytesIO(data)) as source:
        width, height = source.size
    ratio = min(1.0, max_width / max(1, width))
    return Image(BytesIO(data), width=width * ratio, height=height * ratio)


def _build_pdf(request: ExportRequest, mode: str, cache: ImageCache) -> bytes:
    output = BytesIO()
    page_size = A5 if request.settings.paper == "A5" else A4
    margin = 13 * mm if request.settings.paper == "A5" else 17 * mm
    document = BaseDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=request.title,
    )
    column_gap = 7 * mm
    column_count = request.settings.columns
    frame_width = (
        document.width
        if column_count == 1
        else (document.width - column_gap) / 2
    )
    frames = [
        Frame(
            margin + index * (frame_width + column_gap),
            margin,
            frame_width,
            document.height,
            leftPadding=0,
            rightPadding=0,
            topPadding=0,
            bottomPadding=0,
        )
        for index in range(column_count)
    ]
    document.addPageTemplates(PageTemplate(id="content", frames=frames))
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except KeyError:
        pass
    font = "STSong-Light"
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "YantongBody",
        parent=styles["BodyText"],
        fontName=font,
        fontSize=request.settings.fontSize,
        leading=request.settings.fontSize * 1.55,
        spaceAfter=4,
        wordWrap="CJK",
    )
    title_style = ParagraphStyle(
        "YantongTitle",
        parent=body,
        fontSize=request.settings.fontSize + 7,
        leading=request.settings.fontSize + 12,
        alignment=TA_CENTER,
        spaceAfter=10,
    )
    meta_style = ParagraphStyle(
        "YantongMeta",
        parent=body,
        fontSize=max(8, request.settings.fontSize - 2),
        textColor="#66736d",
    )
    story = [_pdf_paragraph(request.title + _suffix(mode), title_style)]
    story.append(_pdf_paragraph(f"共 {len(request.groups)} 个题组", meta_style))
    story.append(HRFlowable(width="100%", thickness=0.7, color="#cfd8d3"))
    story.append(Spacer(1, 6))

    def add_solution(group: ExportGroup, question: ExportQuestion, index: int) -> None:
        label = _question_label(group, question, index)
        story.append(_pdf_paragraph(f"{label}. 答案：{question.answer or '未录入'}", body))
        if request.settings.includeAnalysis and question.analysis:
            story.append(_pdf_paragraph(f"解析：{question.analysis}", body))

    if mode in {"questions", "combined"}:
        for group in request.groups:
            story.append(_pdf_paragraph(f"{group.order}. {group.title}", body))
            if group.meta:
                story.append(_pdf_paragraph(group.meta, meta_style))
            for index, question in enumerate(group.questions):
                label = _question_label(group, question, index)
                story.append(_pdf_paragraph(f"{label}. {question.stem}", body))
                for option in question.options:
                    story.append(_pdf_paragraph(option, body))
            for _ in range(request.settings.blankLines):
                story.append(Spacer(1, request.settings.fontSize * 1.25))
                story.append(HRFlowable(width="100%", thickness=0.35, color="#d9dfdc"))
            if request.settings.packageMode == "combined" and request.settings.solutionMode == "inline":
                for index, question in enumerate(group.questions):
                    add_solution(group, question, index)
            story.append(Spacer(1, max(4, request.settings.questionSpacing / 2)))
    if mode == "solutions" or (
        mode == "combined" and request.settings.solutionMode == "end"
    ):
        if mode == "combined":
            story.append(PageBreak())
            story.append(_pdf_paragraph("答案与解析", title_style))
        for group in request.groups:
            story.append(_pdf_paragraph(f"{group.order}. {group.title}", body))
            for index, question in enumerate(group.questions):
                add_solution(group, question, index)
            story.append(Spacer(1, 6))
    document.build(story)
    data = output.getvalue()
    if len(data) > MAX_OUTPUT_BYTES:
        raise ValueError("生成的 PDF 超过 12MB，请减少题目数量后重试")
    return data


def _docx_font(run, size: int, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold


def _docx_text(document: Document, value: str, size: int, bold: bool = False):
    paragraph = document.add_paragraph()
    _docx_font(paragraph.add_run(str(value or " ")), size, bold)
    return paragraph


def _build_docx(request: ExportRequest, mode: str, cache: ImageCache) -> bytes:
    document = Document()
    section = document.sections[0]
    if request.settings.paper == "A5":
        section.page_width, section.page_height = Mm(148), Mm(210)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(17)
    columns = section._sectPr.xpath("./w:cols")[0]
    columns.set(qn("w:num"), str(request.settings.columns))
    columns.set(qn("w:space"), "360")
    title = document.add_paragraph()
    title.alignment = 1
    _docx_font(title.add_run(request.title + _suffix(mode)), request.settings.fontSize + 7, True)
    _docx_text(document, f"共 {len(request.groups)} 个题组", max(8, request.settings.fontSize - 2))

    def add_picture(image: ExportImage) -> None:
        data = cache.get(image)
        max_width = 112 if request.settings.paper == "A5" else 170
        document.add_picture(BytesIO(data), width=Mm(max_width))

    def add_solution(group: ExportGroup, question: ExportQuestion, index: int) -> None:
        label = _question_label(group, question, index)
        _docx_text(document, f"{label}. 答案：{question.answer or '未录入'}", request.settings.fontSize)
        if request.settings.includeAnalysis and question.analysis:
            _docx_text(document, f"解析：{question.analysis}", request.settings.fontSize)

    if mode in {"questions", "combined"}:
        for group in request.groups:
            _docx_text(document, f"{group.order}. {group.title}", request.settings.fontSize, True)
            if group.meta:
                _docx_text(document, group.meta, max(8, request.settings.fontSize - 2))
            for index, question in enumerate(group.questions):
                label = _question_label(group, question, index)
                _docx_text(document, f"{label}. {question.stem}", request.settings.fontSize)
                for option in question.options:
                    _docx_text(document, option, request.settings.fontSize)
            for _ in range(request.settings.blankLines):
                _docx_text(document, "________________________________________", request.settings.fontSize)
            if request.settings.packageMode == "combined" and request.settings.solutionMode == "inline":
                for index, question in enumerate(group.questions):
                    add_solution(group, question, index)
    if mode == "solutions" or (
        mode == "combined" and request.settings.solutionMode == "end"
    ):
        if mode == "combined":
            document.add_section(WD_SECTION.NEW_PAGE)
            _docx_text(document, "答案与解析", request.settings.fontSize + 5, True)
        for group in request.groups:
            _docx_text(document, f"{group.order}. {group.title}", request.settings.fontSize, True)
            for index, question in enumerate(group.questions):
                add_solution(group, question, index)
    output = BytesIO()
    document.save(output)
    data = output.getvalue()
    if len(data) > MAX_OUTPUT_BYTES:
        raise ValueError("生成的 Word 超过 12MB，请减少题目数量后重试")
    return data


def render_export(request: ExportRequest) -> list[dict[str, object]]:
    cache = ImageCache()
    extension = request.format
    mime_type = (
        "application/pdf"
        if extension == "pdf"
        else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    files = []
    for mode in _document_modes(request):
        data = (
            _build_pdf(request, mode, cache)
            if extension == "pdf"
            else _build_docx(request, mode, cache)
        )
        files.append(
            {
                "name": f"{_safe_name(request.title)}{_suffix(mode)}.{extension}",
                "mimeType": mime_type,
                "size": len(data),
                "contentBase64": base64.b64encode(data).decode("ascii"),
            }
        )
    return files


def _docx_set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.makeelement(qn("w:shd"))
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.first_child_found_in("w:tblW")
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(total))
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(indent_dxa))
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(widths_dxa[index]))
            margins = tc_pr.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def _add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    label = paragraph.add_run("研通专业课练习  ·  ")
    _docx_font(label, 8)
    label.font.color.rgb = RGBColor(117, 127, 141)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    _docx_font(run, 8)
    run.font.color.rgb = RGBColor(117, 127, 141)


def _clozed_text(block, answer_number_start: int) -> tuple[list[tuple[str, bool]], list[dict]]:
    blanks = sorted(
        [item for item in block.blanks if item.enabled], key=lambda item: (item.start, item.end)
    )
    cursor = 0
    runs: list[tuple[str, bool]] = []
    answers: list[dict] = []
    for offset, blank in enumerate(blanks):
        if blank.start < cursor or block.text[blank.start:blank.end] != blank.answer:
            raise ValueError("挖空位置与段落文字不一致，请重新核对后导出")
        if blank.start > cursor:
            runs.append((block.text[cursor:blank.start], False))
        number = answer_number_start + offset
        runs.append((f"〔{number}〕________", True))
        answers.append(
            {
                "number": number,
                "answer": blank.answer,
                "context": block.text[max(0, blank.start - 18):min(len(block.text), blank.end + 18)],
                "page": block.page,
                "block_order": block.order,
            }
        )
        cursor = blank.end
    if cursor < len(block.text):
        runs.append((block.text[cursor:], False))
    return runs, answers


def render_cloze_export(request: ClozeExportRequest) -> list[dict[str, object]]:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(16), Mm(16)
    section.left_margin, section.right_margin = Mm(18), Mm(18)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    _add_page_number(section)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = document.add_paragraph()
    title.alignment = 1
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    _docx_font(title.add_run(request.title), 18, True)
    subtitle = document.add_paragraph()
    subtitle.alignment = 1
    subtitle.paragraph_format.space_after = Pt(12)
    _docx_font(subtitle.add_run(f"{request.courseName} · 填空练习"), 10)
    subtitle.runs[0].font.color.rgb = RGBColor(104, 116, 133)
    identity = document.add_paragraph()
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    identity.paragraph_format.space_after = Pt(16)
    _docx_font(identity.add_run("姓名：____________________    日期：____________________"), 10)

    answers: list[dict] = []
    answer_number = 1
    for block_index, block in enumerate(request.blocks, start=1):
        runs, current_answers = _clozed_text(block, answer_number)
        answer_number += len(current_answers)
        answers.extend(current_answers)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.55
        prefix = paragraph.add_run(f"{block_index}. ")
        _docx_font(prefix, 11, True)
        for value, is_blank in runs:
            run = paragraph.add_run(value)
            _docx_font(run, 11, is_blank)
            if is_blank:
                run.font.color.rgb = RGBColor(70, 87, 168)
    if not answers:
        raise ValueError("请至少保留一处挖空后再生成练习")

    document.add_page_break()
    answer_title = document.add_paragraph()
    answer_title.paragraph_format.space_before = Pt(0)
    answer_title.paragraph_format.space_after = Pt(5)
    _docx_font(answer_title.add_run("答案与核对"), 16, True)
    note = document.add_paragraph()
    _docx_font(note.add_run("请按题号核对答案；上下文用于快速定位原文。"), 10)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths_dxa = [1020, 2460, 6265]
    headers = ["题号", "答案", "原文定位"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        _docx_set_cell_shading(cell, "E8EEF8")
        _docx_font(cell.paragraphs[0].add_run(value), 10, True)
    for answer in answers:
        row = table.add_row().cells
        values = [
            f"{answer['number']}\n第 {answer['block_order']} 段",
            answer["answer"],
            answer["context"],
        ]
        for index, value in enumerate(values):
            paragraph = row[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            _docx_font(paragraph.add_run(value), 9 if index == 2 else 10, index == 1)
    _set_repeat_table_header(table.rows[0])
    _set_table_geometry(table, widths_dxa)

    output = BytesIO()
    document.save(output)
    data = output.getvalue()
    if len(data) > MAX_OUTPUT_BYTES:
        raise ValueError("生成的 Word 超过 12MB，请减少资料内容后重试")
    return [
        {
            "name": f"{_safe_name(request.title)}-填空练习.docx",
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "size": len(data),
            "contentBase64": base64.b64encode(data).decode("ascii"),
        }
    ]
